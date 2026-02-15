"""
Cover letter DOCX builder.
Accepts a cover_letter_text dict (from AI) + profile + job_analysis + options.
Returns a tuple of (io.BytesIO DOCX, preview_text str).
"""
import io
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_font(run, size_pt: float, bold: bool = False, italic: bool = False) -> None:
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic


def _set_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)


def _add_footer_page_number(doc: Document) -> None:
    """Add a centered page number to the footer."""
    section = doc.sections[0]
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run()

    # Insert PAGE field
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    _set_font(run, 10)


def build(
    cover_letter_result: dict,
    profile: dict,
    job_analysis: dict,
    options: dict
) -> tuple[io.BytesIO, str]:
    """
    Build a cover letter DOCX.

    Args:
        cover_letter_result: From GeminiProvider.generate_cover_letter_text() —
            { cover_letter_text: str, subject_line: str }
        profile: MasterProfile dict.
        job_analysis: JobAnalysis dict.
        options: { include_header: bool, include_footer: bool }

    Returns:
        (BytesIO DOCX bytes, preview_text string)
    """
    doc = Document()
    _set_margins(doc)

    # Remove default paragraph spacing
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)

    contact = profile.get('contact', {})
    name = contact.get('name', '')
    email = contact.get('email', '')
    phone = contact.get('phone', '')
    job_title = job_analysis.get('job_title', 'the Position')
    company_name = job_analysis.get('company_name', '')
    letter_text = cover_letter_result.get('cover_letter_text', '')

    # --- Header (optional) ---
    if options.get('include_header', True):
        # Applicant name
        name_para = doc.add_paragraph()
        name_para.paragraph_format.space_after = Pt(2)
        name_run = name_para.add_run(name)
        _set_font(name_run, 12, bold=True)

        # Contact line
        contact_parts = [p for p in [email, phone, contact.get('location')] if p]
        if contact_parts:
            contact_para = doc.add_paragraph()
            contact_para.paragraph_format.space_after = Pt(12)
            _set_font(contact_para.add_run(' | '.join(contact_parts)), 10.5)

        # Date
        date_para = doc.add_paragraph()
        date_para.paragraph_format.space_after = Pt(12)
        _set_font(date_para.add_run(datetime.now().strftime('%B %d, %Y')), 10.5)

        # Company name (if available)
        if company_name:
            co_para = doc.add_paragraph()
            co_para.paragraph_format.space_after = Pt(16)
            _set_font(co_para.add_run(company_name), 10.5)

    # --- Salutation ---
    sal_para = doc.add_paragraph()
    sal_para.paragraph_format.space_after = Pt(12)
    _set_font(sal_para.add_run('Dear Hiring Manager,'), 11)

    # --- Body paragraphs ---
    paragraphs = [p.strip() for p in letter_text.split('\n\n') if p.strip()]
    for i, para_text in enumerate(paragraphs):
        body_para = doc.add_paragraph()
        body_para.paragraph_format.space_after = Pt(10)
        # Remove any leading "Paragraph N:" labels the AI might have included
        cleaned = para_text
        if cleaned.startswith('Paragraph ') and ':' in cleaned[:15]:
            cleaned = cleaned.split(':', 1)[1].strip()
        _set_font(body_para.add_run(cleaned), 11)

    # --- Closing ---
    closing_para = doc.add_paragraph()
    closing_para.paragraph_format.space_before = Pt(6)
    closing_para.paragraph_format.space_after = Pt(24)
    _set_font(closing_para.add_run('Sincerely,'), 11)

    sig_para = doc.add_paragraph()
    sig_para.paragraph_format.space_after = Pt(0)
    _set_font(sig_para.add_run(name), 11, bold=True)

    if email:
        email_para = doc.add_paragraph()
        email_para.paragraph_format.space_after = Pt(0)
        _set_font(email_para.add_run(email), 10.5)

    # --- Footer (optional) ---
    if options.get('include_footer', False):
        _add_footer_page_number(doc)

    # --- Write to BytesIO ---
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # Build preview text
    preview_lines = []
    if options.get('include_header', True):
        preview_lines.append(name)
        preview_lines.append(' | '.join([p for p in [email, phone] if p]))
        preview_lines.append(datetime.now().strftime('%B %d, %Y'))
        if company_name:
            preview_lines.append(company_name)
        preview_lines.append('')
    preview_lines.append('Dear Hiring Manager,')
    preview_lines.append('')
    preview_lines.extend(paragraphs)
    preview_lines.append('')
    preview_lines.append('Sincerely,')
    preview_lines.append(name)

    return buf, '\n'.join(preview_lines)
