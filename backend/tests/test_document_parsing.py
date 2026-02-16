"""
Unit tests for document extraction functions in backend.api.documents.
Tests the three private parsers directly — no Flask, no AI, no network.
"""
import io
import pytest

from backend.api.documents import _extract_txt, _extract_pdf, _extract_docx
from .conftest import fixture_bytes


# ===== TXT extraction ==========================================================

class TestExtractTxt:

    def test_valid_utf8(self):
        result = _extract_txt(b'Hello, world!')
        assert result == 'Hello, world!'

    def test_utf8_accents(self):
        raw = 'résumé: Üniversität Señor'.encode('utf-8')
        result = _extract_txt(raw)
        assert 'résumé' in result
        assert 'Üniversität' in result
        assert 'Señor' in result

    def test_latin1_fallback(self):
        # 'café' encoded as latin-1 (0xe9 for é — not valid UTF-8 alone)
        raw = 'caf\xe9'.encode('latin-1')
        result = _extract_txt(raw)
        assert 'caf' in result  # content preserved via latin-1 fallback

    def test_empty_bytes_returns_empty_string(self):
        result = _extract_txt(b'')
        assert result == ''

    def test_multiline_preserved(self):
        raw = b'Line 1\nLine 2\nLine 3'
        result = _extract_txt(raw)
        assert result == 'Line 1\nLine 2\nLine 3'

    def test_windows_crlf_preserved(self):
        raw = b'Line 1\r\nLine 2\r\n'
        result = _extract_txt(raw)
        assert 'Line 1' in result
        assert 'Line 2' in result

    def test_unicode_cjk_characters(self):
        # Chinese characters — valid UTF-8
        raw = '张三\nSoftware Engineer'.encode('utf-8')
        result = _extract_txt(raw)
        assert '张三' in result

    def test_returns_string_type(self):
        result = _extract_txt(b'test')
        assert isinstance(result, str)


# ===== PDF extraction ==========================================================

class TestExtractPdf:

    def test_valid_pdf_returns_text(self):
        pdf_bytes = fixture_bytes('sample_resume.pdf')
        result = _extract_pdf(pdf_bytes)
        assert isinstance(result, str)
        assert len(result) > 0

    def test_pdf_contains_jane_doe(self):
        pdf_bytes = fixture_bytes('sample_resume.pdf')
        result = _extract_pdf(pdf_bytes)
        assert 'Jane Doe' in result

    def test_pdf_contains_email(self):
        pdf_bytes = fixture_bytes('sample_resume.pdf')
        result = _extract_pdf(pdf_bytes)
        assert 'jane@example.com' in result

    def test_corrupt_pdf_raises(self):
        with pytest.raises(Exception):
            _extract_pdf(b'%PDF-1.4 CORRUPT TRUNCATED DATA')

    def test_empty_bytes_raises_or_empty(self):
        # pdfminer may raise or return empty on 0-byte input — both acceptable
        try:
            result = _extract_pdf(b'')
            assert isinstance(result, str)
        except Exception:
            pass  # pdfminer raising is also valid behaviour


# ===== DOCX extraction =========================================================

class TestExtractDocx:

    def test_valid_docx_returns_text(self):
        docx_bytes = fixture_bytes('sample_resume.docx')
        result = _extract_docx(docx_bytes)
        assert isinstance(result, str)
        assert len(result) > 0

    def test_docx_contains_jane_doe(self):
        docx_bytes = fixture_bytes('sample_resume.docx')
        result = _extract_docx(docx_bytes)
        assert 'Jane Doe' in result

    def test_docx_filters_blank_paragraphs(self):
        """_extract_docx skips paragraphs where para.text.strip() is falsy."""
        from docx import Document
        doc = Document()
        doc.add_paragraph('Jane Doe')
        doc.add_paragraph('')        # blank — must be filtered out
        doc.add_paragraph('   ')     # whitespace-only — must be filtered out
        doc.add_paragraph('Engineer')
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        result = _extract_docx(buf.read())
        non_empty_lines = [l for l in result.split('\n') if l.strip()]
        assert non_empty_lines == ['Jane Doe', 'Engineer']

    def test_docx_all_blank_returns_empty(self):
        """DOCX with only whitespace paragraphs → empty string."""
        from docx import Document
        doc = Document()
        doc.add_paragraph('   ')
        doc.add_paragraph('\t')
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        result = _extract_docx(buf.read())
        assert result == ''

    def test_corrupt_docx_raises(self):
        with pytest.raises(Exception):
            _extract_docx(b'\x00\x01\x02\x03 not a zip file')

    def test_returns_string_type(self):
        docx_bytes = fixture_bytes('sample_resume.docx')
        result = _extract_docx(docx_bytes)
        assert isinstance(result, str)
