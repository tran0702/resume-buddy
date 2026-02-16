"""
scripts/create_test_fixtures.py

One-time script to generate binary test fixtures for the pytest suite.
Run from the project root:
    .venv/Scripts/python scripts/create_test_fixtures.py

Requires: fpdf2 (in requirements-dev.txt)
    .venv/Scripts/pip install -r requirements-dev.txt
"""
import io
import os

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), '..', 'backend', 'tests', 'fixtures')
os.makedirs(FIXTURES_DIR, exist_ok=True)


def create_txt_fixture():
    path = os.path.join(FIXTURES_DIR, 'sample_resume.txt')
    content = (
        'Jane Doe\n'
        'jane@example.com | (555) 123-4567 | New York, NY\n'
        'linkedin.com/in/janedoe | github.com/janedoe\n\n'
        'SUMMARY\n'
        'Experienced software engineer with 5 years building Python APIs and React frontends.\n\n'
        'EXPERIENCE\n'
        'Acme Corporation  |  Software Engineer  |  2022-01 – Present\n'
        '• Built RESTful APIs with Flask and FastAPI serving 50K requests/day\n'
        '• Reduced deployment time by 40% by migrating CI/CD to GitHub Actions\n'
        '• Mentored two junior engineers through code reviews and pair programming\n\n'
        'StartupCo  |  Junior Developer  |  2020-06 – 2021-12\n'
        '• Developed React dashboards for real-time data visualisation\n'
        '• Integrated third-party payment APIs (Stripe, PayPal)\n\n'
        'EDUCATION\n'
        'State University  |  B.S. Computer Science  |  2020-05\n'
        'GPA: 3.8 / 4.0  |  Dean\'s List\n\n'
        'SKILLS\n'
        'Python, JavaScript, TypeScript, React, Flask, FastAPI, PostgreSQL, Docker, Git\n\n'
        'CERTIFICATIONS\n'
        'AWS Certified Developer – Associate (2023)\n'
    )
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)
    print(f'  Created: {path}')


def create_docx_fixture():
    from docx import Document
    path = os.path.join(FIXTURES_DIR, 'sample_resume.docx')
    doc = Document()
    paragraphs = [
        'Jane Doe',
        'jane@example.com | (555) 123-4567 | New York, NY',
        '',
        'SUMMARY',
        'Experienced software engineer with 5 years building Python APIs.',
        '',
        'EXPERIENCE',
        'Acme Corporation — Software Engineer — 2022-01 to Present',
        'Built RESTful APIs with Flask serving 50K requests/day',
        'Reduced deployment time by 40% via GitHub Actions',
        '',
        'EDUCATION',
        'State University — B.S. Computer Science — 2020-05',
        '',
        'SKILLS',
        'Python, JavaScript, React, Flask, PostgreSQL, Docker',
    ]
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(path)
    print(f'  Created: {path}')


def create_pdf_fixture():
    from fpdf import FPDF, XPos, YPos
    path = os.path.join(FIXTURES_DIR, 'sample_resume.pdf')
    pdf = FPDF()
    pdf.add_page()

    # Use ASCII-only strings — Helvetica only supports latin-1
    pdf.set_font('Helvetica', 'B', 16)
    pdf.cell(0, 10, 'Jane Doe', new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')

    pdf.set_font('Helvetica', size=10)
    pdf.cell(0, 6, 'jane@example.com | (555) 123-4567 | New York, NY',
             new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')
    pdf.ln(4)

    pdf.set_font('Helvetica', 'B', 11)
    pdf.cell(0, 8, 'SUMMARY', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', size=10)
    pdf.multi_cell(0, 6,
        'Experienced software engineer with 5 years building Python APIs and React frontends.')
    pdf.ln(2)

    pdf.set_font('Helvetica', 'B', 11)
    pdf.cell(0, 8, 'EXPERIENCE', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 6, 'Acme Corporation - Software Engineer - 2022-01 to Present',
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', size=10)
    pdf.cell(0, 6, '- Built RESTful APIs with Flask serving 50K requests/day',
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(0, 6, '- Reduced deployment time by 40% via GitHub Actions',
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)

    pdf.set_font('Helvetica', 'B', 11)
    pdf.cell(0, 8, 'EDUCATION', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', size=10)
    pdf.cell(0, 6, 'State University - B.S. Computer Science - 2020-05',
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)

    pdf.set_font('Helvetica', 'B', 11)
    pdf.cell(0, 8, 'SKILLS', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', size=10)
    pdf.cell(0, 6, 'Python, JavaScript, React, Flask, PostgreSQL, Docker, Git',
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.output(path)
    print(f'  Created: {path}')


if __name__ == '__main__':
    print('Creating test fixtures...')
    create_txt_fixture()
    create_docx_fixture()
    create_pdf_fixture()
    print('Done.')
